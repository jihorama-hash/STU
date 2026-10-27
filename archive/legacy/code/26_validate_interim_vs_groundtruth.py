from __future__ import annotations

from _paths import BASE

"""중간분석 개정 로직(25번) 결과를 Phase A 지상검증 판정(52건)과 대조한다.

- GT: survey/지상검증_수동조사표_v2.xlsx 4_상충검증 시트 (구 로직 이벤트 + 사람 판정)
- NEW: runs/interim/DJI_0482/pet_events_directional.csv (개정 로직 이벤트)
- 매칭: 시각(±1.0s) + 상충유형 + 상위 ROI(분할 전 기준)
- 산출: GT 각 건의 유지/제거 여부, 정탐률 개선 요약
"""

import csv
from pathlib import Path

from openpyxl import load_workbook

SURVEY = BASE / "survey" / "지상검증_수동조사표_v2.xlsx"
NEW_EVENTS = BASE / "runs" / "interim" / "DJI_0482" / "pet_events_directional.csv"
OUT = BASE / "runs" / "interim" / "DJI_0482" / "groundtruth_comparison.csv"

TIME_TOL = 1.0


def mmss_to_sec(s):
    m, sec = str(s).split(":")
    return int(m) * 60 + float(sec)


def load_gt():
    wb = load_workbook(SURVEY, data_only=True)
    ws = wb["4_상충검증"]
    rows = []
    for r in range(4, 56):
        v = lambda c: ws.cell(row=r, column=c).value
        if v(1) is None:
            continue
        rows.append(
            {
                "no": v(1),
                "t": mmss_to_sec(v(2)),
                "roi": v(3),
                "type": v(4),
                "pet": v(7),
                "critical": v(9) == "예",
                "verdict": (v(11) or "").strip(),
            }
        )
    return rows


def load_new():
    with NEW_EVENTS.open(encoding="utf-8-sig", newline="") as f:
        return [
            {
                "t": float(r["second_entry_sec"]),
                "parent": r["roi_parent"],
                "side": r["roi_side"],
                "type": r["conflict_type"],
                "pet": float(r["pet_sec"]),
                "angle": r["cross_angle_deg"],
            }
            for r in csv.DictReader(f)
        ]


def main():
    gt = load_gt()
    new = load_new()
    used = set()
    out_rows = []
    for g in gt:
        match = None
        for i, n in enumerate(new):
            if i in used:
                continue
            if n["type"] != g["type"] or n["parent"] != g["roi"]:
                continue
            if abs(n["t"] - g["t"]) <= TIME_TOL:
                match = (i, n)
                break
        status = "retained" if match else "removed"
        if match:
            used.add(match[0])
        out_rows.append(
            {
                "gt_no": g["no"],
                "gt_time_sec": round(g["t"], 1),
                "gt_roi": g["roi"],
                "gt_type": g["type"],
                "gt_verdict": g["verdict"],
                "gt_critical": g["critical"],
                "new_status": status,
                "new_side": match[1]["side"] if match else "",
                "new_cross_angle": match[1]["angle"] if match else "",
            }
        )

    with OUT.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(out_rows[0].keys()))
        w.writeheader()
        w.writerows(out_rows)

    def count(rows, verdict, status=None):
        sel = [r for r in rows if r["gt_verdict"] == verdict]
        if status:
            sel = [r for r in sel if r["new_status"] == status]
        return len(sel)

    total_retained = sum(1 for r in out_rows if r["new_status"] == "retained")
    real_kept = count(out_rows, "실제상충", "retained")
    real_total = count(out_rows, "실제상충")
    print(f"GT 52건 중 개정 로직 유지: {total_retained}건")
    for v in ["실제상충", "단순통과", "오탐", "판단불가"]:
        kept = count(out_rows, v, "retained")
        tot = count(out_rows, v)
        print(f"  {v}: {tot}건 중 유지 {kept} / 제거 {tot-kept}")
    if total_retained:
        print(f"유지분 정탐률(실제상충/유지): {real_kept}/{total_retained} = {real_kept/total_retained*100:.1f}%  (구 로직: 6/52 = 11.5%)")
    print(f"실제상충 회수율: {real_kept}/{real_total}")
    unmatched_new = len(new) - len(used)
    print(f"GT와 매칭되지 않은 신규 이벤트: {unmatched_new}건 (분할 ROI로 새로 잡힌 쌍 등)")
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
