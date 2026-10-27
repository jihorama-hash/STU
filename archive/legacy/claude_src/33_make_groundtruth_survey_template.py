"""단일 영상용 지상검증 조사표(xlsx) 생성.

`25_compute_pet_directional.py`가 만든 `pet_events_directional.csv`를 읽어, 다른 영상들의
지상검증 조사표(`survey/지상검증_수동조사표_DJI_0481.xlsx` 등)와 동일한 서식으로 빈 판정란만
남긴 xlsx를 만든다. 자동 산출 가능한 열(이벤트 시각/ROI/유형/PET/교차각/동시체류)만 채우고,
사람이 영상을 보고 채워야 하는 열(critical/판정/심각도/비고)은 비워 둔다.

    python src/33_make_groundtruth_survey_template.py --video DJI_0482
"""

from __future__ import annotations

from _paths import BASE

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PET_DIR = BASE / "runs" / "before_l_mt09"
SURVEY = BASE / "survey"

HEADER = [
    "연번", "이벤트 시각", "ROI", "상충 유형", "선행 객체", "후행 객체",
    "PET(초)", "교차각(°)", "동시체류", "critical", "판정", "심각도(주관)", "비고",
]


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Generate a blank-verdict groundtruth survey xlsx for one video.")
    ap.add_argument("--video", required=True, help="e.g. DJI_0482")
    ap.add_argument(
        "--events",
        type=Path,
        default=None,
        help="pet_events_directional.csv path (default: runs/before_l_mt09/<video>/pet_directional/pet_events_directional.csv)",
    )
    ap.add_argument("--output", type=Path, default=None)
    return ap.parse_args()


def sec_to_mmss(t: float) -> str:
    m = int(t // 60)
    s = t - m * 60
    return f"{m:02d}:{s:04.1f}"


def main() -> None:
    args = parse_args()
    events_path = args.events or PET_DIR / args.video / "pet_directional" / "pet_events_directional.csv"
    output_path = args.output or SURVEY / f"지상검증_수동조사표_{args.video}.xlsx"

    with events_path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    rows.sort(key=lambda r: float(r["second_entry_sec"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "4_상충검증"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADER))
    ws.cell(row=1, column=1, value=(
        f"대상: {args.video} 방향성 PET 이벤트 {len(rows)}건(yolo11l+mt0.9, 최종 ROI + 튜닝된 파라미터: "
        f"교차각40도/PET_pv상한1.5초 기준). 이벤트 시각으로 영상(record/before/{args.video}.MP4) 이동해 "
        "판정. critical·보행자-차량 우선. 판정: 실제상충/단순통과/오탐/판단불가 중 선택."
    ))

    header_fill = PatternFill(start_color="00DDDDDD", end_color="00DDDDDD", fill_type="solid")
    for c, name in enumerate(HEADER, start=1):
        cell = ws.cell(row=3, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for i, r in enumerate(rows, start=1):
        pet = float(r["pet_sec"])
        ws.append(
            [
                i,
                sec_to_mmss(float(r["second_entry_sec"])),
                r["roi_parent"],
                r["conflict_type"],
                r["first_class_name"],
                r["second_class_name"],
                round(pet, 3),
                round(float(r["cross_angle_deg"]), 1),
                "예" if r["overlap"].strip().lower() == "true" else "아니오",
                None,
                None,
                None,
                None,
            ]
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"{len(rows)}건 -> {output_path}")


if __name__ == "__main__":
    main()
