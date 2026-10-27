"""횡단보도 상충 이벤트 지상검증 집계 (Phase B, `25_compute_pet_directional.py` 결과 대상).

단일 영상용 서식(`연번/이벤트시각/ROI/...`)과 다중 영상용 서식(`연번/영상/이벤트시각/ROI/...`)
을 모두 읽어, work plan §11 정의대로 집계한다:

- 정탐률 = 실제상충 / 자동검출 전체 (판정 종류 무관, 자동검출된 모든 이벤트가 모수)
- 회수율 = 자동검출된 실제상충 / (자동검출된 실제상충 + 사람이 영상에서 육안으로 찾아
  "추가"로 적어 넣은 누락 실제상충)

    python src/32_validate_crosswalk_groundtruth.py
"""

from __future__ import annotations

from _paths import BASE

import csv

from openpyxl import load_workbook

SURVEY = BASE / "survey"

# (파일, 단일 영상이면 영상명 고정, 다중 영상 서식이면 None -> '영상' 열에서 읽음)
SURVEY_FILES = [
    (SURVEY / "지상검증_수동조사표_DJI_0481_F.xlsx", "DJI_0481"),
    (SURVEY / "지상검증_수동조사표_DJI_0482_F.xlsx", "DJI_0482"),
    (SURVEY / "지상검증_수동조사표_DJI_0486_Com.xlsx", "DJI_0486"),
    (SURVEY / "지상검증_수동조사표_잔여5편_횡단보도_F.xlsx", None),
]


def read_sheet(path, fixed_video):
    wb = load_workbook(path, data_only=True)
    ws = wb["4_상충검증"]
    header = {
        ws.cell(row=3, column=c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=3, column=c).value
    }
    col_no = header["연번"]
    col_video = header.get("영상")
    col_roi = header["ROI"]
    col_type = header["상충 유형"]
    col_verdict = header["판정"]

    rows = []
    for r in range(4, ws.max_row + 1):
        roi = ws.cell(row=r, column=col_roi).value
        if roi is None:
            continue
        no = ws.cell(row=r, column=col_no).value
        video = fixed_video if fixed_video else ws.cell(row=r, column=col_video).value
        rows.append(
            {
                "video": video,
                "detected": isinstance(no, (int, float)),
                "roi": str(roi),
                "conflict_type": ws.cell(row=r, column=col_type).value,
                "verdict": (ws.cell(row=r, column=col_verdict).value or "").strip(),
            }
        )
    return rows


def summarize(rows, label):
    detected = [r for r in rows if r["detected"]]
    missed_real = [r for r in rows if not r["detected"] and r["verdict"] == "실제상충"]
    total = len(detected)
    real_detected = sum(1 for r in detected if r["verdict"] == "실제상충")
    real_total = real_detected + len(missed_real)
    precision = real_detected / total * 100 if total else float("nan")
    recall = real_detected / real_total * 100 if real_total else float("nan")
    print(
        f"{label}: 자동검출 {total}건 중 실제상충 {real_detected}건 -> 정탐률 {precision:.1f}%"
        f" | 실제상충 총 {real_total}건(누락 {len(missed_real)}) 중 검출 {real_detected}건 -> 회수율 {recall:.1f}%"
    )
    return total, real_detected, real_total


def main():
    all_rows = []
    for path, fixed_video in SURVEY_FILES:
        all_rows.extend(read_sheet(path, fixed_video))
    videos = sorted(set(r["video"] for r in all_rows))

    print(f"=== 영상별 (전체 ROI, {len(videos)}편) ===")
    for v in videos:
        summarize([r for r in all_rows if r["video"] == v], v)

    print()
    print("=== 영상별 (횡단보도만) ===")
    for v in videos:
        vrows = [r for r in all_rows if r["video"] == v and r["roi"].startswith("crosswalk")]
        if vrows:
            summarize(vrows, v)
        else:
            print(f"{v}: 횡단보도 이벤트 없음")

    print()
    print("=== 합계 ===")
    cw_rows = [r for r in all_rows if r["roi"].startswith("crosswalk")]
    summarize(cw_rows, f"횡단보도만 ({len(videos)}편 합계)")
    summarize(all_rows, f"전체 ROI ({len(videos)}편 합계)")

    out_path = BASE / "runs" / "before_l_mt09" / "combined_interim" / "groundtruth_crosswalk_validation.csv"
    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["video", "detected", "roi", "conflict_type", "verdict"])
        w.writeheader()
        w.writerows(all_rows)
    print(f"saved -> {out_path}")


if __name__ == "__main__":
    main()
